import openpyxl, json, os
from datetime import datetime

PATH = r'D:\.goldminer3\projects\fe28bb3f-a453-11ee-b878-14755b767e75\仓位结构.xlsx'
SRC = '可转债等权指数'
wb = openpyxl.load_workbook(PATH, data_only=True)
ws = wb[SRC]

dates, prem, avgp, medp = [], [], [], []
for r in range(2, ws.max_row + 1):
    dt = ws.cell(r, 1).value
    if not isinstance(dt, datetime):
        continue
    dates.append(dt.strftime('%Y-%m-%d'))
    p = ws.cell(r, 11).value
    a = ws.cell(r, 6).value
    m = ws.cell(r, 7).value
    prem.append(round(p * 100, 2) if isinstance(p, (int, float)) else None)
    avgp.append(round(a, 2) if isinstance(a, (int, float)) else None)
    medp.append(round(m, 2) if isinstance(m, (int, float)) else None)


def js_arr(seq):
    return '[' + ','.join('null' if v is None else str(v) for v in seq) + ']'


ASSETS_DIR = r'web\canzhai-trend\assets'

# 输出 data.js（供 <script> 同步加载）
with open(os.path.join(ASSETS_DIR, 'data.js'), 'w', encoding='utf-8') as f:
    f.write('window.CZ_DATA = {\n')
    f.write('  dates: ' + json.dumps(dates) + ',\n')
    f.write('  premium: ' + js_arr(prem) + ',\n')
    f.write('  avgPrice: ' + js_arr(avgp) + ',\n')
    f.write('  medPrice: ' + js_arr(medp) + ',\n')
    f.write('};\n')

# 输出 data.json（供 fetch 异步加载）
with open(os.path.join(ASSETS_DIR, 'data.json'), 'w', encoding='utf-8') as f:
    json.dump({'dates': dates, 'premium': prem, 'avgPrice': avgp, 'medPrice': medp}, f, ensure_ascii=False)

print('rows:', len(dates), dates[0], '->', dates[-1])
print('prem last:', prem[-1], 'avg last:', avgp[-1], 'med last:', medp[-1])